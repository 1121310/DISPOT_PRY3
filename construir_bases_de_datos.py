# -*- coding: utf-8 -*-
"""
Construye las bases de datos (libros Excel) que leen los tres métodos TNEP.

Todas las bases comparten el mismo esquema, de modo que los scripts
tnep_lp.py, tnep_minlp.py y tnep_milp_bigm.py funcionan sin cambios sobre
cualquier sistema:

  Hoja "Parametros" : Parametro | Valor | Descripcion
  Hoja "Barras"     : Barra | Tipo | PG_MW | PL_MW
  Hoja "Lineas"     : Desde | Hasta | Circuitos | x_pu | r_pu | Pmax_MW | Longitud
                      (red existente; x, r y Pmax son POR CIRCUITO)
  Hoja "Corredores" : Corredor | Desde | Hasta | x_pu | r_pu | Pmax_MW |
                      Costo | n_max | Longitud
                      (líneas candidatas; Costo es el costo de UNA línea nueva)
  Hoja "Fuente"     : notas sobre el origen de los datos

Sistemas generados
------------------
1. garver6_tnep.xlsx         Caso base del documento TNEP.pdf (Garver 1970 /
                             Villasana 1984): 6 barras, 6 líneas, 4 corredores.
2. garver6_15corredores.xlsx Garver (1970), Tabla IV completa: los 15 pares
                             de barras como corredores candidatos.
3. sur46_sin_redespacho.xlsx Sistema Sur de Brasil de 46 barras (Romero et
                             al. 2002), sin redespacho.

Uso:  python construir_bases_de_datos.py
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

CARPETA = Path(__file__).resolve().parent

# Resistencia por milla de Garver (1970), Tabla III: R = 0.0025 pu/milla,
# X = 0.01 pu/milla (la misma estructura usada en los Excel del curso).
R_MILLA = 0.0025
X_MILLA = 0.01


# --------------------------------------------------------------------------
# 1) Garver 6 barras: base de datos del documento TNEP.pdf
# --------------------------------------------------------------------------
def garver_documento() -> dict:
    parametros = [
        ("nombre_sistema", "Garver 6 barras - 4 corredores (TNEP.pdf)", "Nombre del caso"),
        ("Sbase_MVA", 100, "Potencia base del sistema"),
        ("barra_referencia", 5, "Barra con angulo cero (TNEP.pdf usa la barra 5)"),
        ("unidad_costo", "pu", "Unidad del costo de inversion"),
        ("bigM_rad", 1.0, "Constante big-M del MILP (Villasana usa 1 rad)"),
        ("bigM_modo", "fijo", "fijo: M igual para todos | auto: camino minimo en la red existente"),
        ("theta_max_rad", 1.5708, "Cota |theta| usada solo para acotar el MINLP (pi/2)"),
        ("criterio_lp", "flujo", "Seleccion del corredor en el LP: flujo | costo"),
        ("tolerancia", 1e-6, "Tolerancia numerica de convergencia"),
        ("mip_gap", 0.0, "Brecha relativa de optimalidad para MILP/MINLP"),
        ("limite_tiempo_s", 3600, "Limite de tiempo por modelo (s)"),
    ]
    # Tabla 1 de TNEP.pdf (despacho economico del anno horizonte)
    barras = [
        (1, "existente", 50, 80),
        (2, "existente", 0, 240),
        (3, "existente", 165, 40),
        (4, "existente", 0, 160),
        (5, "existente", 0, 240),
        (6, "futura", 545, 0),
    ]
    # Tabla 2 de TNEP.pdf (lineas existentes) + R de Garver (1970) Tabla III
    lineas = []
    for de, a, x, pmax, long_mi in [
        (1, 2, 0.4, 100, 40), (1, 4, 0.6, 80, 60), (1, 5, 0.2, 100, 20),
        (2, 3, 0.2, 100, 20), (2, 4, 0.4, 100, 40), (3, 5, 0.2, 100, 20),
    ]:
        lineas.append((de, a, 1, x, round(R_MILLA * long_mi, 4), pmax, long_mi))
    # Tabla 3 de TNEP.pdf (corredores) y seccion 6.2 (x, 100 MW por linea)
    corredores = []
    for de, a, costo, long_mi, nmax in [
        (6, 2, 150, 30, 4), (6, 4, 150, 30, 4), (6, 5, 305, 61, 3), (3, 5, 100, 20, 3),
    ]:
        corredores.append((f"{de}-{a}", de, a, round(X_MILLA * long_mi, 4),
                           round(R_MILLA * long_mi, 4), 100, costo, nmax, long_mi))
    fuente = [
        "Datos tomados de TNEP.pdf (P. M. De Oliveira-De Jesus, 2026): Tablas 1, 2 y 3 y seccion 6.2.",
        "Equivalen a los libros C10_E1/E2/E3 Garver (LP, MILP, MINLP) del repositorio",
        "github.com/pmdeoliveiradejesus/PSE/tree/main/C10_TEP (x = 0.01 pu/milla, Sbase = 100 MVA).",
        "Costo de cada linea nueva = numero guia de Garver (1970) = 5 x longitud en millas.",
        "R = 0.0025 pu/milla (Garver 1970, Tabla III); solo se usa en el post-proceso de perdidas/flujo AC.",
        "Nota: el libro MILP del curso admite hasta 5 lineas en 6-2 y 6-4; TNEP.pdf (Tabla 3) fija 4.",
    ]
    return dict(parametros=parametros, barras=barras, lineas=lineas,
                corredores=corredores, fuente=fuente)


# --------------------------------------------------------------------------
# 2) Garver 6 barras con los 15 corredores de la Tabla IV de Garver (1970)
# --------------------------------------------------------------------------
def garver_15_corredores() -> dict:
    base = garver_documento()
    parametros = [list(p) for p in base["parametros"]]
    for p in parametros:
        if p[0] == "nombre_sistema":
            p[1] = "Garver 6 barras - 15 corredores (Garver 1970, Tabla IV)"
    # Garver (1970), Tabla IV: terminales, millas, numero guia de la trayectoria
    # de sobrecarga (costo), R, X y capacidad del circuito nuevo.
    tabla_iv = [
        (1, 2, 40, 200, 0.10, 0.40, 100), (1, 3, 38, 190, 0.09, 0.38, 100),
        (1, 4, 60, 300, 0.15, 0.60, 80), (1, 5, 20, 100, 0.05, 0.20, 100),
        (1, 6, 68, 340, 0.17, 0.68, 70), (2, 3, 20, 100, 0.05, 0.20, 100),
        (2, 4, 40, 200, 0.10, 0.40, 100), (2, 5, 31, 155, 0.08, 0.31, 100),
        (2, 6, 30, 150, 0.08, 0.30, 100), (3, 4, 59, 295, 0.15, 0.59, 82),
        (3, 5, 20, 100, 0.05, 0.20, 100), (3, 6, 48, 240, 0.12, 0.48, 100),
        (4, 5, 63, 315, 0.16, 0.63, 75), (4, 6, 30, 150, 0.08, 0.30, 100),
        (5, 6, 61, 305, 0.15, 0.61, 78),
    ]
    corredores = [(f"{de}-{a}", de, a, x, r, pmax, costo, 5, millas)
                  for de, a, millas, costo, r, x, pmax in tabla_iv]
    fuente = [
        "Barras y red existente iguales a garver6_tnep.xlsx (TNEP.pdf, Tablas 1 y 2).",
        "Corredores: Garver (1970), Tabla IV (overload-path and new-circuit data), 15 pares de barras.",
        "Costo = numero guia de la trayectoria de sobrecarga (5 x millas), igual que en TNEP.pdf.",
        "Garver (1970) no limita el numero de circuitos por corredor; aqui se supone n_max = 5.",
    ]
    return dict(parametros=parametros, barras=base["barras"], lineas=base["lineas"],
                corredores=corredores, fuente=fuente)


# --------------------------------------------------------------------------
# 3) Sistema Sur de Brasil, 46 barras, caso sin redespacho de generacion
# --------------------------------------------------------------------------
# Red existente: (desde, hasta, circuitos n0, x por circuito, r, Pmax por circuito MW)
SUR46_EXISTENTES = [
    (1, 7, 1, 0.0616, 0.0121, 270), (1, 2, 2, 0.1065, 0.0209, 270),
    (4, 9, 1, 0.0924, 0.0181, 270), (5, 9, 1, 0.1173, 0.0230, 270),
    (5, 8, 1, 0.1132, 0.0222, 270), (7, 8, 1, 0.1023, 0.0200, 270),
    (4, 5, 2, 0.0566, 0.0111, 270), (2, 5, 2, 0.0324, 0.0064, 270),
    (8, 13, 1, 0.1348, 0.0264, 240), (9, 14, 2, 0.1756, 0.0344, 220),
    (12, 14, 2, 0.0740, 0.0145, 270), (14, 18, 2, 0.1514, 0.0297, 240),
    (13, 18, 1, 0.1805, 0.0354, 220), (13, 20, 1, 0.1073, 0.0210, 270),
    (18, 20, 1, 0.1997, 0.0391, 200), (19, 21, 1, 0.0278, 0.0019, 1500),
    (16, 17, 1, 0.0078, 0.0005, 2000), (17, 19, 1, 0.0061, 0.0004, 2000),
    (14, 26, 1, 0.1614, 0.0316, 220), (14, 22, 1, 0.0840, 0.0165, 270),
    (22, 26, 1, 0.0790, 0.0155, 270), (20, 23, 2, 0.0932, 0.0183, 270),
    (23, 24, 2, 0.0774, 0.0152, 270), (26, 27, 2, 0.0832, 0.0163, 270),
    (24, 34, 1, 0.1647, 0.0323, 220), (24, 33, 1, 0.1448, 0.0284, 240),
    (33, 34, 1, 0.1265, 0.0248, 270), (27, 36, 1, 0.0915, 0.0179, 270),
    (27, 38, 2, 0.2080, 0.0407, 200), (36, 37, 1, 0.1057, 0.0207, 270),
    (34, 35, 2, 0.0491, 0.0096, 270), (35, 38, 1, 0.1980, 0.0388, 200),
    (37, 39, 1, 0.0283, 0.0055, 270), (37, 40, 1, 0.1281, 0.0251, 270),
    (37, 42, 1, 0.2105, 0.0412, 200), (39, 42, 3, 0.2030, 0.0398, 200),
    (40, 42, 1, 0.0932, 0.0183, 270), (38, 42, 3, 0.0907, 0.0178, 270),
    (32, 43, 1, 0.0309, 0.0021, 1400), (42, 44, 1, 0.1206, 0.0236, 270),
    (44, 45, 1, 0.1864, 0.0365, 200), (19, 32, 1, 0.0195, 0.0013, 1800),
    (46, 19, 1, 0.0222, 0.0015, 1800), (46, 16, 1, 0.0203, 0.0014, 1800),
    (18, 19, 1, 0.0125, 0.0, 600), (20, 21, 1, 0.0125, 0.0, 600),
    (42, 43, 1, 0.0125, 0.0, 600),
]
# Corredores candidatos: (desde, hasta, x, r, Pmax MW, costo MUS$ por linea)
SUR46_CORREDORES = [
    (20, 23, 0.0932, 0.0183, 270, 6.268), (33, 34, 0.1265, 0.0248, 270, 8.288),
    (42, 44, 0.1206, 0.0236, 270, 7.934), (20, 21, 0.0125, 0.0, 600, 8.178),
    (42, 43, 0.0125, 0.0, 600, 8.178), (46, 6, 0.0128, 0.0009, 2000, 16.005),
    (19, 25, 0.0325, 0.0022, 1400, 37.748), (24, 25, 0.0125, 0.0, 600, 8.178),
    (5, 6, 0.0125, 0.0, 600, 8.178), (1, 7, 0.0616, 0.0121, 270, 4.349),
    (1, 2, 0.1065, 0.0209, 270, 7.076), (4, 9, 0.0924, 0.0181, 270, 6.217),
    (5, 9, 0.1173, 0.0230, 270, 7.732), (5, 8, 0.1132, 0.0222, 270, 7.480),
    (7, 8, 0.1023, 0.0200, 270, 6.823), (4, 5, 0.0566, 0.0111, 270, 4.046),
    (2, 5, 0.0324, 0.0064, 270, 2.581), (8, 13, 0.1348, 0.0264, 240, 8.793),
    (9, 14, 0.1756, 0.0344, 220, 11.267), (12, 14, 0.0740, 0.0145, 270, 5.106),
    (14, 18, 0.1514, 0.0297, 240, 9.803), (13, 18, 0.1805, 0.0354, 220, 11.570),
    (13, 20, 0.1073, 0.0210, 270, 7.126), (18, 20, 0.1997, 0.0391, 200, 12.732),
    (19, 21, 0.0278, 0.0019, 1500, 32.632), (16, 17, 0.0078, 0.0005, 2000, 10.505),
    (17, 19, 0.0061, 0.0004, 2000, 8.715), (14, 26, 0.1614, 0.0316, 220, 10.409),
    (14, 22, 0.0840, 0.0165, 270, 5.712), (22, 26, 0.0790, 0.0155, 270, 5.409),
    (23, 24, 0.0774, 0.0152, 270, 5.308), (26, 27, 0.0832, 0.0163, 270, 5.662),
    (24, 34, 0.1647, 0.0323, 220, 10.611), (24, 33, 0.1448, 0.0284, 240, 9.399),
    (27, 36, 0.0915, 0.0179, 270, 6.167), (27, 38, 0.2080, 0.0407, 200, 13.237),
    (36, 37, 0.1057, 0.0207, 270, 7.025), (34, 35, 0.0491, 0.0096, 270, 3.591),
    (35, 38, 0.1980, 0.0388, 200, 12.631), (37, 39, 0.0283, 0.0055, 270, 2.329),
    (37, 40, 0.1281, 0.0251, 270, 8.389), (37, 42, 0.2105, 0.0412, 200, 13.388),
    (39, 42, 0.2030, 0.0398, 200, 12.934), (40, 42, 0.0932, 0.0183, 270, 6.268),
    (38, 42, 0.0907, 0.0178, 270, 6.116), (32, 43, 0.0309, 0.0021, 1400, 35.957),
    (44, 45, 0.1864, 0.0365, 200, 11.924), (19, 32, 0.0195, 0.0013, 1800, 23.423),
    (46, 19, 0.0222, 0.0015, 1800, 26.365), (46, 16, 0.0203, 0.0014, 1800, 24.319),
    (18, 19, 0.0125, 0.0, 600, 8.178), (2, 4, 0.0882, 0.0173, 270, 5.965),
    (14, 15, 0.0374, 0.0073, 270, 2.884), (46, 10, 0.0081, 0.0006, 2000, 10.889),
    (4, 11, 0.2246, 0.0440, 240, 14.247), (5, 11, 0.0915, 0.0179, 270, 6.167),
    (46, 3, 0.0203, 0.0014, 1800, 24.319), (16, 28, 0.0222, 0.0015, 1800, 26.365),
    (16, 32, 0.0311, 0.0021, 1400, 36.213), (17, 32, 0.0232, 0.0016, 1700, 27.516),
    (21, 25, 0.0174, 0.0012, 2000, 21.121), (25, 32, 0.0319, 0.0022, 1400, 37.109),
    (31, 32, 0.0046, 0.0003, 2000, 7.052), (28, 31, 0.0053, 0.0004, 2000, 7.819),
    (28, 30, 0.0058, 0.0004, 2000, 8.331), (27, 29, 0.0998, 0.0196, 270, 6.672),
    (26, 29, 0.0541, 0.0106, 270, 3.894), (28, 41, 0.0339, 0.0023, 1300, 39.283),
    (28, 43, 0.0406, 0.0028, 1200, 46.701), (31, 41, 0.0278, 0.0019, 1500, 32.632),
    (32, 41, 0.0309, 0.0021, 1400, 35.957), (41, 43, 0.0139, 0.0010, 2000, 17.284),
    (40, 45, 0.2205, 0.0432, 180, 13.994), (15, 16, 0.0125, 0.0, 600, 8.178),
    (46, 11, 0.0125, 0.0, 600, 8.178), (29, 30, 0.0125, 0.0, 600, 8.178),
    (40, 41, 0.0125, 0.0, 600, 8.178), (2, 3, 0.0125, 0.0, 600, 8.178),
    (9, 10, 0.0125, 0.0, 600, 8.178),
]
# Demanda (MW) y generacion sin redespacho (MW); suman 6880 MW.
SUR46_CARGA = {2: 443.1, 4: 300.7, 5: 238.0, 8: 72.2, 12: 511.9, 13: 185.8,
               20: 1091.2, 22: 81.9, 23: 458.1, 24: 478.2, 26: 231.9, 33: 229.1,
               35: 216.0, 36: 90.1, 38: 216.0, 40: 262.1, 42: 1607.9, 44: 79.1,
               45: 86.7}
SUR46_GENERACION = {14: 944, 16: 1366, 17: 1000, 19: 773, 27: 54, 28: 730,
                    31: 310, 32: 450, 34: 221, 37: 212, 39: 221, 46: 599}


def sur46(n_max: int = 3) -> dict:
    parametros = [
        ("nombre_sistema", "Sur de Brasil 46 barras - sin redespacho", "Nombre del caso"),
        ("Sbase_MVA", 100, "Potencia base del sistema"),
        ("barra_referencia", 16, "Barra con angulo cero"),
        ("unidad_costo", "MUS$", "Unidad del costo de inversion (millones de US$)"),
        ("bigM_rad", 10.0, "Constante big-M del MILP (rad); ver analisis de sensibilidad"),
        ("bigM_modo", "auto", "fijo: M igual para todos | auto: camino minimo en la red existente"),
        ("theta_max_rad", 3.1416, "Cota |theta| usada solo para acotar el MINLP (pi)"),
        ("criterio_lp", "flujo", "Seleccion del corredor en el LP: flujo | costo"),
        ("tolerancia", 1e-6, "Tolerancia numerica de convergencia"),
        ("mip_gap", 0.0, "Brecha relativa de optimalidad para MILP/MINLP"),
        ("limite_tiempo_s", 3600, "Limite de tiempo por modelo (s)"),
    ]
    con_red = {b for de, a, *_ in SUR46_EXISTENTES for b in (de, a)}
    barras = [(b, "existente" if b in con_red else "nueva",
               SUR46_GENERACION.get(b, 0), SUR46_CARGA.get(b, 0.0)) for b in range(1, 47)]
    lineas = [(de, a, n0, x, r, pmax, None) for de, a, n0, x, r, pmax in SUR46_EXISTENTES]
    corredores = [(f"{de}-{a}", de, a, x, r, pmax, costo, n_max, None)
                  for de, a, x, r, pmax, costo in SUR46_CORREDORES]
    fuente = [
        "Sistema Sur de Brasil de 46 barras, 62 circuitos existentes y 79 corredores, demanda 6880 MW.",
        "Sistema de prueba descrito en Romero et al. (2002), IEE Proc. Gener. Transm. Distrib. 149(1):27-36.",
        "Datos digitales tomados del repositorio github.com/MahdiMehrtash/Global-TEP (Data46) y",
        "convertidos al modelo DC clasico: capacidad/1.1 y circuitos paralelos existentes separados (n0).",
        "Verificados fila por fila contra S. Taglialenha, tesis doctoral UNESP (2008), Apendice A",
        "(Tablas A.3-A.4): 78 de 79 corredores coinciden; se corrigio el costo del corredor 23-24",
        "(5.308 MUS$; la fuente digital trae 35.308). Generacion sin redespacho segun la misma tesis.",
        "Optimo DC sin redespacho: 154.420 MUS$ (Romero y Monticelli, 1994; optimalidad probada por",
        "Binato, Pereira y Granville, 2001). Costo en millones de US$ por linea.",
    ]
    return dict(parametros=parametros, barras=barras, lineas=lineas,
                corredores=corredores, fuente=fuente)


# --------------------------------------------------------------------------
def escribir_libro(datos: dict, ruta: Path) -> None:
    """Escribe el diccionario de datos en un libro Excel con formato legible."""
    hojas = {
        "Parametros": pd.DataFrame(datos["parametros"], columns=["Parametro", "Valor", "Descripcion"]),
        "Barras": pd.DataFrame(datos["barras"], columns=["Barra", "Tipo", "PG_MW", "PL_MW"]),
        "Lineas": pd.DataFrame(datos["lineas"], columns=["Desde", "Hasta", "Circuitos", "x_pu",
                                                         "r_pu", "Pmax_MW", "Longitud"]),
        "Corredores": pd.DataFrame(datos["corredores"],
                                   columns=["Corredor", "Desde", "Hasta", "x_pu", "r_pu",
                                            "Pmax_MW", "Costo", "n_max", "Longitud"]),
        "Fuente": pd.DataFrame({"Notas": datos["fuente"]}),
    }
    with pd.ExcelWriter(ruta, engine="openpyxl") as escritor:
        for nombre, df in hojas.items():
            df.to_excel(escritor, sheet_name=nombre, index=False)
    libro = load_workbook(ruta)
    relleno = PatternFill("solid", fgColor="DDE7F3")
    for hoja in libro.worksheets:
        for celda in hoja[1]:
            celda.font = Font(bold=True)
            celda.fill = relleno
            celda.alignment = Alignment(horizontal="center")
        for columna in hoja.columns:
            ancho = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
            hoja.column_dimensions[columna[0].column_letter].width = min(max(10, ancho + 2), 95)
        hoja.freeze_panes = "A2"
    libro.save(ruta)


def main() -> None:
    casos = {
        "garver6_tnep.xlsx": garver_documento(),
        "garver6_15corredores.xlsx": garver_15_corredores(),
        "sur46_sin_redespacho.xlsx": sur46(),
    }
    for nombre, datos in casos.items():
        ruta = CARPETA / nombre
        escribir_libro(datos, ruta)
        print(f"Base de datos escrita: {ruta.name}  "
              f"({len(datos['barras'])} barras, {len(datos['lineas'])} lineas, "
              f"{len(datos['corredores'])} corredores)")


if __name__ == "__main__":
    main()
